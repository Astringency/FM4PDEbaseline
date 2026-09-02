#!/usr/bin/env python
"""Build a compact multi-sheet Main/ID/Rough evaluation workbook.

The workbook is built directly from the completed ``summary.json`` files for
the selected rows in ``main_results.jsonl``.  Each task group is written to a
separate sheet and only the applicable relative-L2 errors for ``a`` and ``u``
are included.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BASELINES = ("fno", "deeponet", "ifno", "recfno", "senseiver", "voronoicnn")
PDE_ORDER = {"poisson": 0, "helmholtz": 1, "darcy": 2, "nsnonbounded": 3, "burger": 4}
BASELINE_ORDER = {name: index for index, name in enumerate(DEFAULT_BASELINES)}
PDE_LABEL = {
    "poisson": "Poisson",
    "helmholtz": "Helmholtz",
    "darcy": "Darcy",
    "nsnonbounded": "NS",
    "burger": "Burgers",
}
BASELINE_LABEL = {
    "fno": "FNO",
    "deeponet": "DeepONet",
    "ifno": "IFNO",
    "recfno": "RecFNO",
    "senseiver": "Senseiver",
    "voronoicnn": "VoronoiCNN",
}
SHEET_DEFINITIONS = (
    ("Full Forward", "full_forward_main", "u"),
    ("Full Inverse", "full_inverse_main", "a"),
    ("Sparse Inverse", "sparse_inverse_main_amortized", "a"),
    ("Sparse Solution", "sparse_solution_main_amortized", "au"),
    ("Burger Time Slices", "sparse_solution_burger_time_slices", "au"),
    ("Sparse Forward", "sparse_forward_main_amortized", "u"),
)


def detect_output_root() -> Path:
    configured = os.environ.get("FM_OUTPUT_ROOT", "").strip()
    if configured:
        return Path(configured).expanduser().resolve()
    candidates = (
        Path.home() / "share/outputs/FM4PDEbaseline",
        Path.home() / "share/zhangxfA100/large_storage/outputs/FM4PDEbaseline",
    )
    return next((path.resolve() for path in candidates if path.is_dir()), candidates[0])


def split_selection(value: str | Iterable[str] | None) -> list[str]:
    if value is None:
        return []
    values = re.split(r"[\s,]+", value.strip()) if isinstance(value, str) else list(value)
    selected: list[str] = []
    for item in values:
        normalized = str(item).strip().lower()
        if normalized and normalized not in selected:
            selected.append(normalized)
    return selected


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    output_root = detect_output_root()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-root",
        type=Path,
        default=output_root,
        help="FM4PDEbaseline output root (default: FM_OUTPUT_ROOT or detected shared output root)",
    )
    parser.add_argument(
        "--matrix",
        type=Path,
        default=None,
        help="main_results.jsonl path (default: OUTPUT_ROOT/matrices/main_results.jsonl)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="output workbook path (default: OUTPUT_ROOT/FM4PDEbaseline_Main_ID_Rough_results.xlsx)",
    )
    parser.add_argument(
        "--baseline-list",
        default=os.environ.get("BASELINE_LIST", ",".join(DEFAULT_BASELINES)),
        help="comma/space-separated baselines (default: BASELINE_LIST or six reusable methods)",
    )
    return parser.parse_args(argv)


def read_json(path: Path) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise RuntimeError(f"summary missing: {path}") from exc
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"summary is not valid JSON: {path}: {exc}") from exc
    if not isinstance(payload, dict):
        raise RuntimeError(f"summary is not a JSON object: {path}")
    return payload


def load_matrix(path: Path, baselines: set[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except FileNotFoundError as exc:
        raise RuntimeError(f"matrix missing: {path}") from exc
    for line_number, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"invalid matrix JSON at line {line_number}: {exc}") from exc
        if row.get("baseline") in baselines and not row.get("skip_reason"):
            rows.append(row)
    if not rows:
        raise RuntimeError(f"matrix has no runnable rows for baselines: {sorted(baselines)}")
    return rows


def summary_path(output_root: Path, row: dict[str, Any], distribution: str) -> Path:
    if distribution == "main":
        namespace = output_root / "runs/main_results"
        run_dir = f"run={row['run_id']}"
    else:
        namespace = output_root / "runs/evaluations" / distribution
        run_dir = f"run=eval_{distribution}_{row['run_id']}"
    return (
        namespace
        / f"task_group={row['task_group']}"
        / f"pde={row['pde']}"
        / f"baseline={row['baseline']}"
        / f"seed={row['seed']}"
        / run_dir
        / "summary.json"
    )


def validated_summary(
    output_root: Path,
    row: dict[str, Any],
    distribution: str,
) -> dict[str, Any]:
    path = summary_path(output_root, row, distribution)
    summary = read_json(path)
    expected_run_id = row["run_id"] if distribution == "main" else f"eval_{distribution}_{row['run_id']}"
    expected = {
        "run_id": expected_run_id,
        "task_group": row["task_group"],
        "pde": row["pde"],
        "baseline": row["baseline"],
        "seed": row["seed"],
        "status": "success",
    }
    for field, value in expected.items():
        if summary.get(field) != value:
            raise RuntimeError(
                f"summary validation failed for {path}: {field}={summary.get(field)!r}, expected {value!r}"
            )
    if int(summary.get("test_size", 0) or 0) != 1000:
        raise RuntimeError(f"summary validation failed for {path}: test_size must be 1000")
    return summary


def metric(summary: dict[str, Any], field: str, context: str) -> float:
    key = f"relative_l2_{field}_mean"
    value = summary.get(key)
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"missing/non-numeric {key} for {context}: {value!r}") from exc
    if not math.isfinite(number):
        raise RuntimeError(f"non-finite {key} for {context}: {number!r}")
    return number * 100.0


def collect_records(
    output_root: Path,
    matrix: Path,
    baselines: set[str],
) -> list[dict[str, Any]]:
    group_order = {group: index for index, (_, group, _) in enumerate(SHEET_DEFINITIONS)}
    records: list[dict[str, Any]] = []
    for row in load_matrix(matrix, baselines):
        if row.get("task_group") not in group_order:
            raise RuntimeError(f"unsupported task group in selected matrix row: {row.get('task_group')!r}")
        records.append(
            {
                "matrix": row,
                "main": validated_summary(output_root, row, "main"),
                "id": validated_summary(output_root, row, "id"),
                "rough": validated_summary(output_root, row, "rough"),
            }
        )
    records.sort(
        key=lambda item: (
            group_order[item["matrix"]["task_group"]],
            PDE_ORDER[item["matrix"]["pde"]],
            BASELINE_ORDER.get(item["matrix"]["baseline"], 999),
            item["matrix"]["baseline"],
            item["matrix"]["seed"],
        )
    )
    return records


def add_summary_sheet(
    workbook: Workbook,
    records: list[dict[str, Any]],
    output_root: Path,
    matrix: Path,
) -> None:
    ws = workbook.active
    ws.title = "Summary"
    navy, blue, light_blue, white, dark = "17365D", "2F75B5", "D9EAF7", "FFFFFF", "1F1F1F"
    thin_gray = Side(style="thin", color="D9E1F2")
    ws.merge_cells("A1:D1")
    ws["A1"] = "FM4PDEbaseline — Main / ID / Rough Evaluation Results"
    ws["A1"].font = Font(size=16, bold=True, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    notes = (
        ("Generated", datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M:%S %Z")),
        ("Units", "Mean relative L2 error (%) — lower is better"),
        ("Main definition", "Original main_results; treated as Smooth for Poisson/Helmholtz/Darcy/NS"),
        ("Burgers caveat", "Burgers Main is the legacy unsuffixed/ID-style test set, not a new smooth evaluation"),
        ("Evaluation protocol", "seed=1, test_size=1000 per run, eval_only using trained checkpoints"),
        ("Metric mapping", "a = relative L2 input/coefficient error; u = relative L2 solution error"),
        ("Source matrix", str(matrix)),
        ("Main source", str(output_root / "runs/main_results")),
        ("ID source", str(output_root / "runs/evaluations/id")),
        ("Rough source", str(output_root / "runs/evaluations/rough")),
    )
    for row_number, (label, value) in enumerate(notes, start=3):
        ws.cell(row_number, 1, label)
        ws.cell(row_number, 2, value)
        ws.cell(row_number, 1).font = Font(bold=True, color=dark)
        ws.cell(row_number, 1).fill = PatternFill("solid", fgColor=light_blue)
        ws.cell(row_number, 1).border = Border(bottom=thin_gray)
        ws.cell(row_number, 2).border = Border(bottom=thin_gray)
        ws.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row_number, start_column=2, end_row=row_number, end_column=4)
    start = len(notes) + 5
    for column, header in enumerate(("Sheet", "Task group", "Rows", "Metrics"), start=1):
        cell = ws.cell(start, column, header)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.alignment = Alignment(horizontal="center")
    for offset, (sheet_name, group, kind) in enumerate(SHEET_DEFINITIONS, start=1):
        count = sum(item["matrix"]["task_group"] == group for item in records)
        for column, value in enumerate(
            (sheet_name, group, count, "a and u" if kind == "au" else kind),
            start=1,
        ):
            ws.cell(start + offset, column, value)
            ws.cell(start + offset, column).border = Border(bottom=thin_gray)
        ws.cell(start + offset, 1).hyperlink = f"#'{sheet_name}'!A1"
        ws.cell(start + offset, 1).style = "Hyperlink"
    ws.freeze_panes = "A3"
    for column, width in {"A": 24, "B": 55, "C": 12, "D": 16}.items():
        ws.column_dimensions[column].width = width
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def task_values(item: dict[str, Any], kind: str) -> list[Any]:
    row = item["matrix"]
    values: list[Any] = [PDE_LABEL[row["pde"]], BASELINE_LABEL.get(row["baseline"], row["baseline"])]
    context = f"{row['task_group']}/{row['pde']}/{row['baseline']}/seed={row['seed']}"
    if kind == "au":
        for field in ("input_or_coeff", "solution"):
            values.extend(metric(item[distribution], field, context) for distribution in ("main", "id", "rough"))
    else:
        field = "input_or_coeff" if kind == "a" else "solution"
        values.extend(metric(item[distribution], field, context) for distribution in ("main", "id", "rough"))
    return values


def add_task_sheets(workbook: Workbook, records: list[dict[str, Any]]) -> dict[str, int]:
    navy, blue, light_gray, white = "17365D", "2F75B5", "F2F2F2", "FFFFFF"
    thin_gray = Side(style="thin", color="D9E1F2")
    row_counts: dict[str, int] = {}
    for sheet_name, group, kind in SHEET_DEFINITIONS:
        ws = workbook.create_sheet(sheet_name)
        subset = [item for item in records if item["matrix"]["task_group"] == group]
        row_counts[sheet_name] = len(subset)
        headers = (
            ["PDE", "Method", "a Main (%)", "a ID (%)", "a Rough (%)", "u Main (%)", "u ID (%)", "u Rough (%)"]
            if kind == "au"
            else ["PDE", "Method", f"{kind} Main (%)", f"{kind} ID (%)", f"{kind} Rough (%)"]
        )
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(1, 1, f"{sheet_name} — Relative L2 Error")
        ws.cell(1, 1).font = Font(size=15, bold=True, color=white)
        ws.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
        ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        ws.cell(2, 1, "Mean over 1,000 test samples; values are percentages; lower is better. Main = original main_results.")
        ws.cell(2, 1).font = Font(italic=True, color="595959")
        header_row = 4
        for column, header in enumerate(headers, start=1):
            cell = ws.cell(header_row, column, header)
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        data_start = header_row + 1
        for row_number, item in enumerate(subset, start=data_start):
            values = task_values(item, kind)
            pde = item["matrix"]["pde"]
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row_number, column, value)
                cell.border = Border(bottom=thin_gray)
                if column >= 3:
                    cell.number_format = "0.000"
                    cell.alignment = Alignment(horizontal="right")
                if PDE_ORDER[pde] % 2 == 1:
                    cell.fill = PatternFill("solid", fgColor=light_gray)
        data_end = data_start + len(subset) - 1
        last_column = chr(64 + len(headers))
        if subset:
            table = Table(
                displayName="Results" + "".join(character for character in sheet_name if character.isalnum()),
                ref=f"A{header_row}:{last_column}{data_end}",
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=False,
                showColumnStripes=False,
            )
            ws.add_table(table)
            ranges = ((3, 5), (6, 8)) if kind == "au" else ((3, 5),)
            for start_column, end_column in ranges:
                ws.conditional_formatting.add(
                    f"{chr(64 + start_column)}{data_start}:{chr(64 + end_column)}{data_end}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="E2F0D9",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFF2CC",
                        end_type="max",
                        end_color="F4CCCC",
                    ),
                )
        ws.freeze_panes = f"A{data_start}"
        ws.auto_filter.ref = f"A{header_row}:{last_column}{data_end}"
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 16
        for column in range(3, len(headers) + 1):
            ws.column_dimensions[chr(64 + column)].width = 15
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"1:{header_row}"
    return row_counts


def validate_workbook(
    path: Path,
    records: list[dict[str, Any]],
    row_counts: dict[str, int],
) -> None:
    with zipfile.ZipFile(path) as archive:
        corrupt_member = archive.testzip()
        if corrupt_member:
            raise RuntimeError(f"generated workbook has a corrupt ZIP member: {corrupt_member}")
    workbook = load_workbook(path, data_only=False)
    expected_sheets = ["Summary", *(definition[0] for definition in SHEET_DEFINITIONS)]
    if workbook.sheetnames != expected_sheets:
        raise RuntimeError(f"generated workbook sheets differ: {workbook.sheetnames}")
    for sheet_name, group, kind in SHEET_DEFINITIONS:
        ws = workbook[sheet_name]
        subset = [item for item in records if item["matrix"]["task_group"] == group]
        if ws.max_row != 4 + row_counts[sheet_name]:
            raise RuntimeError(f"generated workbook row count differs for {sheet_name}")
        for row_number, item in enumerate(subset, start=5):
            expected = task_values(item, kind)
            observed = [ws.cell(row_number, column).value for column in range(1, len(expected) + 1)]
            for column, (actual, wanted) in enumerate(zip(observed, expected), start=1):
                if isinstance(wanted, float):
                    if not math.isclose(float(actual), wanted, rel_tol=1e-13, abs_tol=1e-13):
                        raise RuntimeError(f"generated workbook value differs at {sheet_name}!{row_number},{column}")
                elif actual != wanted:
                    raise RuntimeError(f"generated workbook value differs at {sheet_name}!{row_number},{column}")
    workbook.close()


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    output_root = args.output_root.expanduser().resolve()
    matrix = (args.matrix or output_root / "matrices/main_results.jsonl").expanduser().resolve()
    output = (args.output or output_root / "FM4PDEbaseline_Main_ID_Rough_results.xlsx").expanduser().resolve()
    baselines = set(split_selection(args.baseline_list))
    unknown = baselines - set(DEFAULT_BASELINES)
    if unknown:
        raise RuntimeError(f"unsupported baseline(s): {sorted(unknown)}")
    records = collect_records(output_root, matrix, baselines)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.stem}.tmp{output.suffix}")
    try:
        workbook = Workbook()
        add_summary_sheet(workbook, records, output_root, matrix)
        row_counts = add_task_sheets(workbook, records)
        workbook.save(temporary)
        workbook.close()
        validate_workbook(temporary, records, row_counts)
        temporary.replace(output)
    finally:
        if temporary.exists():
            temporary.unlink()
    print(
        json.dumps(
            {
                "output": str(output),
                "matrix": str(matrix),
                "selected_baselines": sorted(baselines),
                "result_rows": len(records),
                "sheet_rows": row_counts,
                "validation": "passed",
            },
            indent=2,
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
