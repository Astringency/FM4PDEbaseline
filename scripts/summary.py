#!/usr/bin/env python
"""Collect main and ablation evaluations into one compact XLSX workbook.

For main-result runs, explicit evaluations under ``runs/evaluations`` are the
preferred source for all three distributions.  When an explicit Smooth result
is missing, spatial-PDE metrics fall back to the historical Smooth evaluation
stored with training under ``runs/main_results``.  Burgers is excluded from
that fallback because its historical unsuffixed test set is ID-like.

Successful ablation evaluations stored below an ``ablation=*`` directory in
``runs/evaluations/{smooth,id,rough}`` are appended to the same workbook.
The workbook contains task settings and metric means with their stored
standard deviations; unavailable metrics or statistics are left blank.
"""

from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


SUMMARY_COLUMNS = [
    "PDE",
    "Method",
    "TASK",
    "DIST",
    "CONDITION",
    "SENSOR",
    "SEED",
    "Ablation",
    "rel L2(a)",
    "rel L2(a) std",
    "rel L2(u)",
    "rel L2(u) std",
    "joint rel L2",
    "joint rel L2 std",
    "pde L",
    "pde L std",
]

_DISTRIBUTIONS = ("smooth", "id", "rough")
_DISTRIBUTION_ORDER = {value: index for index, value in enumerate(_DISTRIBUTIONS)}
_CONDITION_ORDER = {"a_only": 0, "u_only": 1, "both": 2}
_IGNORED_RESULT_DIRECTORIES = {"archive", "historical", "quarantine"}
_SMOOTH_FALLBACK_PDES = {"poisson", "helmholtz", "darcy", "nsnonbounded"}
_METRIC_FIELDS = {
    "a": "relative_l2_input_or_coeff",
    "u": "relative_l2_solution",
}
_MULTICONDITION_METRIC_FIELDS = {
    "a": "rel_l2_a",
    "u": "rel_l2_u",
}
_SUMMARY_IDENTITY_FIELDS = ("task_group", "pde", "baseline", "seed")
_TASK_LABELS = {
    "forward": "forward",
    "sparse_forward": "forward",
    "inverse": "inverse",
    "sparse_inverse": "inverse",
    "sparse_solution": "both",
    "sparse_solution_multicondition": "both",
}
_PDE_LABELS = {
    "poisson": "Poisson",
    "helmholtz": "Helmholtz",
    "darcy": "Darcy",
    "burger": "Burgers",
    "nsnonbounded": "NS",
}
_METHOD_LABELS = {
    "fno": "FNO",
    "deeponet": "DeepONet",
    "ifno": "IFNO",
    "recfno": "RecFNO",
    "senseiver": "Senseiver",
    "voronoicnn": "VoronoiCNN",
    "pinn_sparse": "PINN-Sparse",
    "pde_opt": "PDE-Opt",
    "pc_bnn": "PC-BNN",
    "var4d": "Var4D",
    "vivid": "VIVID",
}


def _resolve_out_root(out_root: str | Path | None) -> Path:
    if out_root is None:
        configured = os.environ.get("OUT_ROOT", "").strip()
        if not configured:
            configured = os.environ.get("FM_OUTPUT_ROOT", "").strip()
        out_root = configured or "outputs"
    return Path(out_root).expanduser().resolve()


def _read_summary(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"invalid summary JSON: {path}: {exc}") from exc
    except OSError as exc:
        raise RuntimeError(f"cannot read summary: {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise RuntimeError(f"summary is not a JSON object: {path}")
    return value


def _load_matrix(root: Path) -> list[dict[str, Any]]:
    path = root / "matrices" / "main_results.jsonl"
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except OSError as exc:
        raise RuntimeError(f"cannot read main-results matrix: {path}: {exc}") from exc
    rows: list[dict[str, Any]] = []
    for line_number, line in enumerate(lines, start=1):
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"invalid matrix JSON at {path}:{line_number}: {exc}") from exc
        if not isinstance(row, dict):
            raise RuntimeError(f"matrix row is not a JSON object at {path}:{line_number}")
        if not row.get("skip_reason"):
            rows.append(row)
    if not rows:
        raise RuntimeError(f"main-results matrix has no runnable rows: {path}")
    return rows


def _main_summary_path(root: Path, row: Mapping[str, Any]) -> Path:
    configured = str(row.get("output_dir", ""))
    marker = "/runs/main_results/"
    if marker in configured:
        relative = configured.split(marker, 1)[1]
        return root / "runs" / "main_results" / relative / "summary.json"
    if configured:
        return Path(configured).expanduser() / "summary.json"
    return (
        root
        / "runs"
        / "main_results"
        / f"task_group={row['task_group']}"
        / f"pde={row['pde']}"
        / f"baseline={row['baseline']}"
        / f"seed={row['seed']}"
        / f"run={row['run_id']}"
        / "summary.json"
    )


def _evaluation_summary_path(
    root: Path,
    row: Mapping[str, Any],
    distribution: str,
) -> Path:
    return (
        root
        / "runs"
        / "evaluations"
        / distribution
        / f"task_group={row['task_group']}"
        / f"pde={row['pde']}"
        / f"baseline={row['baseline']}"
        / f"seed={row['seed']}"
        / f"run=eval_{distribution}_{row['run_id']}"
        / "summary.json"
    )


def _load_result(
    path: Path,
    row: Mapping[str, Any],
    *,
    expected_run_id: str,
    required: bool,
) -> dict[str, Any] | None:
    if not path.is_file():
        if required:
            raise RuntimeError(f"main-results summary is missing: {path}")
        return None
    item = _read_summary(path)
    if item.get("status") != "success":
        if required:
            raise RuntimeError(f"main-results summary is not successful: {path}")
        return None
    expected = {field: row.get(field) for field in _SUMMARY_IDENTITY_FIELDS}
    expected["run_id"] = expected_run_id
    for field, value in expected.items():
        if item.get(field) != value:
            raise RuntimeError(
                f"summary identity mismatch at {path}: "
                f"{field}={item.get(field)!r}, expected {value!r}"
            )
    return item


def _relative_l2_stats(
    item: Mapping[str, Any] | None, metric: str
) -> tuple[float | str, float | str]:
    """Read the mean and standard deviation from the same metric family."""
    if item is None:
        return "", ""
    fields = [_METRIC_FIELDS[metric]]
    if item.get("task") == "sparse_solution_multicondition":
        fields.insert(0, _MULTICONDITION_METRIC_FIELDS[metric])
    for field in fields:
        mean = _finite_value(item, f"{field}_mean")
        if mean != "":
            return mean, _finite_value(item, f"{field}_std")
    return "", ""


def _finite_value(item: Mapping[str, Any] | None, field: str) -> float | str:
    if item is None:
        return ""
    try:
        value = float(item.get(field))
    except (TypeError, ValueError):
        return ""
    return value if math.isfinite(value) else ""


def _task_label(row: Mapping[str, Any]) -> str:
    task = str(row.get("task", ""))
    try:
        return _TASK_LABELS[task]
    except KeyError as exc:
        raise RuntimeError(f"unsupported task in main-results matrix: {task!r}") from exc


def _sensor_label(item: Mapping[str, Any]) -> str:
    mode = str(item.get("sensor_mode", item.get("requested_sensor_mode", ""))).lower()
    if mode in {"", "none"}:
        return ""
    if "time_slice" in mode or "sensor_col" in mode or "sersor_col" in mode:
        return "sersor_col"
    return "random"


def _display_label(value: Any, labels: Mapping[str, str]) -> str:
    text = str(value)
    return labels.get(text.lower(), text)


def _distribution_label(distribution: str) -> str:
    return "ID" if distribution == "id" else distribution.title()


def _result_row(
    *,
    identity: Mapping[str, Any],
    distribution: str,
    result: Mapping[str, Any] | None,
    sensor_fallback: Mapping[str, Any] | None,
    ablation: str,
    condition: str,
) -> dict[str, Any]:
    task = _task_label(identity)
    rel_a, rel_a_std = _relative_l2_stats(result, "a")
    rel_u, rel_u_std = _relative_l2_stats(result, "u")
    return {
        "PDE": _display_label(identity.get("pde", ""), _PDE_LABELS),
        "Method": _display_label(identity.get("baseline", ""), _METHOD_LABELS),
        "TASK": task,
        "DIST": _distribution_label(distribution),
        "CONDITION": condition,
        "SENSOR": _sensor_label(result or sensor_fallback or {}),
        "SEED": identity.get("seed", ""),
        "Ablation": ablation,
        "rel L2(a)": rel_a,
        "rel L2(a) std": rel_a_std,
        "rel L2(u)": rel_u,
        "rel L2(u) std": rel_u_std,
        "joint rel L2": _finite_value(result, "joint_rel_l2_mean"),
        "joint rel L2 std": _finite_value(result, "joint_rel_l2_std"),
        "pde L": _finite_value(result, "pde_residual_mean") if task == "both" else "",
        "pde L std": _finite_value(result, "pde_residual_std") if task == "both" else "",
    }


def _evaluation_ablation_rows(root: Path) -> list[dict[str, Any]]:
    rows: list[tuple[tuple[Any, ...], dict[str, Any]]] = []
    evaluation_root = root / "runs" / "evaluations"
    for distribution in _DISTRIBUTIONS:
        distribution_root = evaluation_root / distribution
        if not distribution_root.is_dir():
            continue
        for summary_path in distribution_root.glob("ablation=*/**/summary.json"):
            relative_parts = summary_path.relative_to(distribution_root).parts
            if any(
                part.lower() in _IGNORED_RESULT_DIRECTORIES
                for part in relative_parts
            ):
                continue
            ablation_parts = [
                part for part in relative_parts if part.startswith("ablation=")
            ]
            if not ablation_parts:
                continue
            result = _read_summary(summary_path)
            if result.get("status") != "success":
                continue
            task = str(result.get("task", ""))
            if task not in _TASK_LABELS:
                raise RuntimeError(
                    f"unsupported task in evaluation summary: {summary_path}: {task!r}"
                )
            ablation = ablation_parts[0].split("=", 1)[1]
            condition = str(
                result.get("evaluation_condition_mode")
                or result.get("condition_mode")
                or ""
            )
            item = _result_row(
                identity=result,
                distribution=distribution,
                result=result,
                sensor_fallback=None,
                ablation=ablation,
                condition=condition,
            )
            sort_key = (
                _DISTRIBUTION_ORDER[distribution],
                ablation,
                str(result.get("pde", "")),
                str(result.get("baseline", "")),
                _CONDITION_ORDER.get(condition, len(_CONDITION_ORDER)),
                str(result.get("run_id", "")),
            )
            rows.append((sort_key, item))
    rows.sort(key=lambda pair: pair[0])
    return [item for _, item in rows]


def collect_summary_rows(out_root: str | Path | None = None) -> list[dict[str, Any]]:
    """Return main-result distribution rows plus explicit ablation evaluations."""

    root = _resolve_out_root(out_root)
    rows: list[dict[str, Any]] = []
    for row in _load_matrix(root):
        run_id = str(row["run_id"])
        main_path = _main_summary_path(root, row)
        main = _load_result(
            main_path,
            row,
            expected_run_id=run_id,
            required=True,
        )
        assert main is not None
        for distribution in _DISTRIBUTIONS:
            evaluation_path = _evaluation_summary_path(root, row, distribution)
            result = _load_result(
                evaluation_path,
                row,
                expected_run_id=f"eval_{distribution}_{run_id}",
                required=False,
            )
            if (
                distribution == "smooth"
                and result is None
                and str(row.get("pde", "")).lower() in _SMOOTH_FALLBACK_PDES
            ):
                result = main
            rows.append(
                _result_row(
                    identity=row,
                    distribution=distribution,
                    result=result,
                    sensor_fallback=main,
                    ablation="",
                    condition="",
                )
            )
    rows.extend(_evaluation_ablation_rows(root))
    return rows


def summary(out_root: str | Path | None = None) -> Path:
    """Write ``OUT_ROOT/summary/results.xlsx`` and return its path."""

    root = _resolve_out_root(out_root)
    rows = collect_summary_rows(root)
    output_dir = root / "summary"
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / "results.xlsx"
    temporary = output_dir / ".results.tmp.xlsx"
    try:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Results"
        worksheet.append(SUMMARY_COLUMNS)
        for row in rows:
            worksheet.append([row[column] for column in SUMMARY_COLUMNS])

        header_fill = PatternFill("solid", fgColor="2F75B5")
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        worksheet.freeze_panes = "A2"
        last_column = get_column_letter(len(SUMMARY_COLUMNS))
        worksheet.auto_filter.ref = f"A1:{last_column}{len(rows) + 1}"
        setting_widths = {
            "PDE": 14,
            "Method": 16,
            "TASK": 10,
            "DIST": 10,
            "CONDITION": 14,
            "SENSOR": 12,
            "SEED": 8,
            "Ablation": 34,
        }
        for index, column in enumerate(SUMMARY_COLUMNS, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = (
                setting_widths.get(column, max(14, len(column) + 2))
            )
            if column in setting_widths:
                continue
            number_format = "0.000000E+00" if column.startswith("pde L") else "0.000000"
            for row_number in range(2, len(rows) + 2):
                worksheet.cell(row_number, index).number_format = number_format
        table = Table(
            displayName="ResultsTable",
            ref=f"A1:{last_column}{len(rows) + 1}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
        workbook.save(temporary)
        workbook.close()

        validation = load_workbook(temporary, read_only=True, data_only=True)
        try:
            if validation.sheetnames != ["Results"]:
                raise RuntimeError(f"unexpected workbook sheets: {validation.sheetnames}")
            saved = validation["Results"]
            headers = [
                saved.cell(1, column).value
                for column in range(1, len(SUMMARY_COLUMNS) + 1)
            ]
            if headers != SUMMARY_COLUMNS or saved.max_row != len(rows) + 1:
                raise RuntimeError("saved workbook structure does not match the summary rows")
        finally:
            validation.close()
        temporary.replace(output)
        legacy_csv = output_dir / "results.csv"
        if legacy_csv.is_file():
            legacy_csv.unlink()
    finally:
        if temporary.exists():
            temporary.unlink()
    return output


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out-root",
        "--output-root",
        dest="out_root",
        default=None,
        help="experiment output root (default: OUT_ROOT, FM_OUTPUT_ROOT, or outputs)",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    output = summary(args.out_root)
    workbook = load_workbook(output, read_only=True)
    try:
        rows = workbook["Results"].max_row - 1
    finally:
        workbook.close()
    print(json.dumps({"output": str(output), "rows": rows}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
