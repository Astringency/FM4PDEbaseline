from __future__ import annotations

import json
import os
from types import SimpleNamespace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts import build_experiment_matrix, collect_results, plot_results, run_experiments


def _write_matrix(path: Path, rows: list[dict]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(row) + "\n" for row in rows), encoding="utf-8")
    return path


def test_matrix_cli_defaults_to_formal_main_results(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    captured: list[str] = []
    monkeypatch.setattr(build_experiment_matrix, "build_main", lambda argv: captured.extend(argv))

    assert build_experiment_matrix.main(
        ["--output-root", str(tmp_path), "--data-manifest", "report.json"]
    ) == 0

    assert captured == [
        "--config",
        "configs/experiments/main_results.yaml",
        "--output-root",
        str(tmp_path),
        "--matrix-name",
        "main_results",
        "--data-manifest",
        "report.json",
    ]


def test_runner_selects_pending_rows_and_assigns_gpu_slots(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    done = tmp_path / "done"
    done.mkdir()
    (done / "run.done").touch()
    matrix = _write_matrix(
        tmp_path / "matrices" / "main.jsonl",
        [
            {"run_id": "done", "output_dir": str(done)},
            {"run_id": "a", "output_dir": str(tmp_path / "a")},
            {"run_id": "b", "output_dir": str(tmp_path / "b")},
            {"run_id": "c", "output_dir": str(tmp_path / "c")},
        ],
    )
    launched: list[tuple[int, str]] = []

    def fake_run(matrix_path: Path, index: int, gpu: str, data_root: Path) -> int:
        launched.append((index, gpu))
        return 0

    monkeypatch.setattr(run_experiments, "run_index", fake_run)
    monkeypatch.setattr(run_experiments, "is_complete", lambda row: row["run_id"] == "done")
    rc = run_experiments.main(
        [str(matrix), "--data-root", str(tmp_path), "--gpus", "2,5", "--jobs-per-gpu", "2"]
    )

    assert rc == 0
    assert sorted(launched) == [(1, "2"), (2, "2"), (3, "5")]


def test_runner_does_not_duplicate_an_active_row(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    running = tmp_path / "running"
    running.mkdir()
    (running / "run.running").touch()
    matrix = _write_matrix(
        tmp_path / "matrix.jsonl",
        [{"run_id": "active", "output_dir": str(running)}],
    )
    launched: list[int] = []
    monkeypatch.setattr(
        run_experiments,
        "run_index",
        lambda _matrix, index, _gpu, _data_root: launched.append(index) or 0,
    )

    assert run_experiments.main([str(matrix), "--data-root", str(tmp_path)]) == 0
    assert launched == []


def test_runner_recovery_restarts_only_dead_running_processes(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
):
    active = tmp_path / "active"
    stale = tmp_path / "stale"
    active.mkdir()
    stale.mkdir()
    (active / "run.running").write_text(
        json.dumps(
            {
                "pid": os.getpid(),
                "pid_start_ticks": run_experiments.process_start_ticks(os.getpid()),
            }
        ),
        encoding="utf-8",
    )
    (stale / "run.running").write_text(json.dumps({"pid": 999_999_999}), encoding="utf-8")
    matrix = _write_matrix(
        tmp_path / "matrix.jsonl",
        [
            {"run_id": "active", "output_dir": str(active)},
            {"run_id": "stale", "output_dir": str(stale)},
        ],
    )
    launched: list[int] = []
    monkeypatch.setattr(
        run_experiments,
        "run_index",
        lambda _matrix, index, _gpu, _data_root: launched.append(index) or 0,
    )
    monkeypatch.setattr(run_experiments, "quarantine_invalid_output", lambda _row: None)

    assert run_experiments.main(
        [str(matrix), "--data-root", str(tmp_path), "--rerun-running"]
    ) == 0
    assert launched == [1]


def test_runner_recovery_rejects_a_reused_pid_identity(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
):
    output = tmp_path / "reused"
    output.mkdir()
    actual_ticks = run_experiments.process_start_ticks(os.getpid())
    (output / "run.running").write_text(
        json.dumps({"pid": os.getpid(), "pid_start_ticks": actual_ticks + 1}), encoding="utf-8"
    )
    matrix = _write_matrix(
        tmp_path / "matrix.jsonl",
        [{"run_id": "reused", "output_dir": str(output)}],
    )
    launched: list[int] = []
    monkeypatch.setattr(
        run_experiments,
        "run_index",
        lambda _matrix, index, _gpu, _data_root: launched.append(index) or 0,
    )
    monkeypatch.setattr(run_experiments, "quarantine_invalid_output", lambda _row: None)

    assert run_experiments.main(
        [str(matrix), "--data-root", str(tmp_path), "--rerun-running"]
    ) == 0
    assert launched == [0]


def test_runner_recovery_treats_a_malformed_process_identity_as_stale(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
):
    output = tmp_path / "malformed"
    output.mkdir()
    (output / "run.running").write_text(
        json.dumps({"pid": os.getpid(), "pid_start_ticks": "not-an-integer"}), encoding="utf-8"
    )
    matrix = _write_matrix(
        tmp_path / "matrix.jsonl",
        [{"run_id": "malformed", "output_dir": str(output)}],
    )
    launched: list[int] = []
    monkeypatch.setattr(
        run_experiments,
        "run_index",
        lambda _matrix, index, _gpu, _data_root: launched.append(index) or 0,
    )
    monkeypatch.setattr(run_experiments, "quarantine_invalid_output", lambda _row: None)

    assert run_experiments.main(
        [str(matrix), "--data-root", str(tmp_path), "--rerun-running"]
    ) == 0
    assert launched == [0]


def test_runner_dry_run_supports_first_and_index_ranges(tmp_path: Path, capsys: pytest.CaptureFixture[str]):
    rows = [{"run_id": str(i), "output_dir": str(tmp_path / str(i))} for i in range(6)]
    matrix = _write_matrix(tmp_path / "matrix.jsonl", rows)

    rc = run_experiments.main(
        [str(matrix), "--data-root", str(tmp_path), "--gpus", "0", "--indices", "1,3-5", "--first", "2", "--dry-run"]
    )

    assert rc == 0
    report = json.loads(capsys.readouterr().out)
    assert [row["index"] for row in report["selected"]] == [1, 3]


def test_runner_waits_for_unresolved_checkpoint_dependencies(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
):
    rows = [
        {
            "run_id": "waiting",
            "output_dir": str(tmp_path / "waiting"),
            "dependency_pending": True,
            "dependency_run_id": "forward",
        },
        {"run_id": "ready", "output_dir": str(tmp_path / "ready")},
    ]
    matrix = _write_matrix(tmp_path / "matrix.jsonl", rows)

    rc = run_experiments.main(
        [str(matrix), "--data-root", str(tmp_path), "--gpus", "0", "--dry-run"]
    )

    assert rc == 0
    report = json.loads(capsys.readouterr().out)
    assert report["waiting"] == 1
    assert [item["run_id"] for item in report["selected"]] == ["ready"]


def test_collect_results_writes_compact_summary_and_allows_execution_cohorts(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]
):
    run_dirs = [tmp_path / "runs" / "train", tmp_path / "runs" / "eval"]
    for run_dir, mode, rel_a, rel_u in (
        (run_dirs[0], "train", None, 0.1),
        (run_dirs[1], "eval_only", 0.2, None),
    ):
        run_dir.mkdir(parents=True)
        (run_dir / "summary.json").write_text(
            json.dumps(
                {
                    "task_group": "full_forward_main" if mode == "train" else "full_inverse_main",
                    "pde": "poisson",
                    "task": "forward" if mode == "train" else "inverse",
                    "baseline": "ifno",
                    "seed": 1,
                    "execution_mode": mode,
                    "test_size": 1000,
                    "relative_l2_input_or_coeff_mean": rel_a,
                    "relative_l2_input_or_coeff_std": 0.01 if rel_a is not None else None,
                    "relative_l2_input_or_coeff_ci95": 0.001 if rel_a is not None else None,
                    "relative_l2_input_or_coeff_n": 1000 if rel_a is not None else 0,
                    "relative_l2_solution_mean": rel_u,
                    "relative_l2_solution_std": 0.01 if rel_u is not None else None,
                    "relative_l2_solution_ci95": 0.001 if rel_u is not None else None,
                    "relative_l2_solution_n": 1000 if rel_u is not None else 0,
                }
            ),
            encoding="utf-8",
        )
    matrix = _write_matrix(
        tmp_path / "matrices" / "main_results.jsonl",
        [
            {"run_id": "train", "output_dir": str(run_dirs[0])},
            {"run_id": "eval", "output_dir": str(run_dirs[1])},
        ],
    )
    monkeypatch.setattr(
        collect_results,
        "collect_validated_results",
        lambda _rows: SimpleNamespace(
            records=[
                {"run_id": "train", "cohort_id": "training"},
                {"run_id": "eval", "cohort_id": "evaluation"},
            ],
            quarantine=[],
            missing_run_ids=[],
        ),
    )
    out = tmp_path / "collected"
    assert collect_results.main(
        [str(matrix), "--output-dir", str(out), "--latex"]
    ) == 0

    rows = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    assert len(rows) == 2
    assert list(rows[0]) == collect_results.COMPACT_COLUMNS
    forward = next(row for row in rows if row["execution_mode"] == "train")
    inverse = next(row for row in rows if row["execution_mode"] == "eval_only")
    assert forward["relative_l2_a_mean"] == ""
    assert forward["relative_l2_u_mean"] == pytest.approx(0.1)
    assert inverse["relative_l2_a_mean"] == pytest.approx(0.2)
    assert inverse["relative_l2_u_mean"] == ""
    assert (out / "summary.csv").is_file()
    assert (out / "latex_table.tex").is_file()
    workbook = load_workbook(out / "results.xlsx", read_only=True, data_only=True)
    assert workbook.sheetnames == ["results", "manifest"]
    assert workbook["results"].max_row == 3
    assert workbook["results"].max_column == len(collect_results.COMPACT_COLUMNS)
    report = json.loads(capsys.readouterr().out)
    assert report["cohort_count"] == 2
    assert report["cohort_counts"] == {"evaluation": 1, "training": 1}


def test_plot_results_renders_from_matrix_manifest(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    run_dir = tmp_path / "run"
    samples = run_dir / "samples"
    samples.mkdir(parents=True)
    manifest = samples / "manifest.jsonl"
    manifest.write_text("{}\n", encoding="utf-8")
    (run_dir / "summary.json").write_text(
        json.dumps({"sample_manifest_path": str(manifest)}), encoding="utf-8"
    )
    matrix = _write_matrix(
        tmp_path / "matrix.jsonl",
        [{"run_id": "r1", "output_dir": str(run_dir)}],
    )
    calls: list[tuple[Path, Path, bool, int]] = []

    def fake_render(source: Path, output_dir: Path, *, force: bool, max_samples: int) -> dict:
        calls.append((Path(source), Path(output_dir), force, max_samples))
        return {"sample_count": 1, "rendered_count": 1, "skipped_count": 0, "errors": []}

    monkeypatch.setattr(plot_results, "_render_manifest", fake_render)

    assert plot_results.main(["--matrix", str(matrix), "--max-samples", "100"]) == 0
    assert calls == [(manifest, run_dir / "samples_pdf", False, 100)]


def test_plot_results_limits_each_manifest_to_first_n_samples(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
):
    manifest = _write_matrix(
        tmp_path / "samples" / "manifest.jsonl",
        [
            {"sample_ordinal": index, "artifact_path": str(tmp_path / f"sample_{index}.pt")}
            for index in range(3)
        ],
    )

    def fake_sample_pdf(_artifact: Path, output: Path, **_kwargs) -> Path:
        Path(output).write_bytes(b"%PDF\n%%EOF\n")
        return Path(output)

    monkeypatch.setattr(plot_results, "render_evaluation_sample_pdf", fake_sample_pdf)
    output = tmp_path / "pdfs"
    report = plot_results._render_manifest(
        manifest,
        output,
        force=False,
        max_samples=2,
    )

    assert report["source_sample_count"] == 3
    assert report["sample_count"] == 2
    assert sorted(path.name for path in output.glob("sample_*.pdf")) == [
        "sample_000000.pdf",
        "sample_000001.pdf",
    ]


def test_collect_results_refuses_incomplete_matrix_before_writing_tables(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
):
    run_dir = tmp_path / "missing"
    matrix = _write_matrix(
        tmp_path / "matrices" / "main_results.jsonl",
        [{"run_id": "missing", "output_dir": str(run_dir)}],
    )
    assert collect_results.main([str(matrix), "--output-dir", str(tmp_path / "aggregate")]) == 2
    assert not (tmp_path / "aggregate").exists()


def test_plot_results_reports_missing_sample_manifest(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "results_raw.jsonl").write_text("{}\n", encoding="utf-8")
    (run_dir / "summary.json").write_text(
        json.dumps({"sample_manifest_path": str(run_dir / "missing.jsonl")}), encoding="utf-8"
    )
    matrix = _write_matrix(
        tmp_path / "matrix.jsonl",
        [{"run_id": "r1", "output_dir": str(run_dir)}],
    )
    assert plot_results.main(["--matrix", str(matrix)]) == 2
