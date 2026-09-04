from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts import summary as result_summary


def _write_result(
    root: Path,
    *,
    run_id: str,
    task_group: str,
    baseline: str,
    rel_a: float | None,
    rel_u: float | None,
    task: str = "sparse_inverse",
    sensor_mode: str = "random_per_sample",
    pde_loss: float | None = None,
    pde: str = "poisson",
    test_file: str | None = None,
    extra: dict | None = None,
) -> None:
    root.mkdir(parents=True, exist_ok=True)
    payload = {
        "status": "success",
        "run_id": run_id,
        "task_group": task_group,
        "pde": pde,
        "task": task,
        "baseline": baseline,
        "seed": 1,
        "sensor_mode": sensor_mode,
        "relative_l2_input_or_coeff_mean": rel_a,
        "relative_l2_solution_mean": rel_u,
        "pde_residual_mean": pde_loss,
    }
    if test_file is not None:
        payload["data_files_json"] = json.dumps({"test": test_file})
    if extra:
        payload.update(extra)
    (root / "summary.json").write_text(
        json.dumps(payload),
        encoding="utf-8",
    )


def _write_matrix(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(row) + "\n" for row in rows), encoding="utf-8")


def test_summary_keeps_all_main_results_and_leaves_missing_evaluations_blank(
    tmp_path: Path,
) -> None:
    main_root = tmp_path / "runs" / "main_results"
    eval_root = tmp_path / "runs" / "evaluations"
    trained_dir = main_root / "trained"
    physics_dir = main_root / "physics"
    _write_matrix(
        tmp_path / "matrices" / "main_results.jsonl",
        [
            {
                "run_id": "trained",
                "task_group": "sparse_solution_main_amortized",
                "pde": "poisson",
                "task": "sparse_solution",
                "baseline": "recfno",
                "seed": 1,
                "output_dir": str(trained_dir),
            },
            {
                "run_id": "physics",
                "task_group": "sparse_inverse_main",
                "pde": "poisson",
                "task": "sparse_inverse",
                "baseline": "pde_opt",
                "seed": 1,
                "output_dir": str(physics_dir),
            },
        ],
    )
    _write_result(
        trained_dir,
        run_id="trained",
        task_group="sparse_solution_main_amortized",
        baseline="recfno",
        rel_a=0.1,
        rel_u=0.2,
        task="sparse_solution",
        pde_loss=0.01,
    )
    _write_result(
        physics_dir,
        run_id="physics",
        task_group="sparse_inverse_main",
        baseline="pde_opt",
        rel_a=0.3,
        rel_u=None,
        pde_loss=0.9,
    )
    _write_result(
        eval_root
        / "id/task_group=sparse_solution_main_amortized/pde=poisson"
        / "baseline=recfno/seed=1/run=eval_id_trained",
        run_id="eval_id_trained",
        task_group="sparse_solution_main_amortized",
        baseline="recfno",
        rel_a=0.4,
        rel_u=0.5,
        task="sparse_solution",
        pde_loss=0.02,
    )
    _write_result(
        eval_root
        / "rough/task_group=sparse_solution_main_amortized/pde=poisson"
        / "baseline=recfno/seed=1/run=eval_rough_trained",
        run_id="eval_rough_trained",
        task_group="sparse_solution_main_amortized",
        baseline="recfno",
        rel_a=0.6,
        rel_u=0.7,
        task="sparse_solution",
        sensor_mode="time_slices_per_sample",
        pde_loss=0.03,
    )

    summary_dir = tmp_path / "summary"
    summary_dir.mkdir()
    (summary_dir / "results.csv").write_text("legacy", encoding="utf-8")

    output = result_summary.summary(tmp_path)

    assert output == tmp_path / "summary" / "results.xlsx"
    assert {path.name for path in output.parent.iterdir()} == {"results.xlsx"}
    workbook = load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == ["Results"]
    worksheet = workbook["Results"]
    assert [cell.value for cell in worksheet[1]] == result_summary.SUMMARY_COLUMNS
    rows = [
        dict(zip(result_summary.SUMMARY_COLUMNS, values))
        for values in worksheet.iter_rows(min_row=2, values_only=True)
    ]
    workbook.close()
    assert len(rows) == 6
    trained_smooth = next(
        row for row in rows if row["Method"] == "RecFNO" and row["DIST"] == "Smooth"
    )
    trained_rough = next(
        row for row in rows if row["Method"] == "RecFNO" and row["DIST"] == "Rough"
    )
    physics_smooth = next(
        row for row in rows if row["Method"] == "PDE-Opt" and row["DIST"] == "Smooth"
    )
    physics_id = next(
        row for row in rows if row["Method"] == "PDE-Opt" and row["DIST"] == "ID"
    )
    assert trained_smooth["TASK"] == "both"
    assert trained_smooth["SENSOR"] == "random"
    assert trained_smooth["rel L2(a)"] == pytest.approx(0.1)
    assert trained_smooth["rel L2(u)"] == pytest.approx(0.2)
    assert trained_smooth["pde L"] == pytest.approx(0.01)
    assert trained_smooth["Namespace"] == "main_results"
    assert trained_smooth["Result Source"] == "main_legacy_smooth_fallback"
    assert trained_smooth["Fallback Used"] is True
    assert trained_rough["SENSOR"] == "sersor_col"
    assert trained_rough["rel L2(a)"] == pytest.approx(0.6)
    assert physics_smooth["TASK"] == "inverse"
    assert physics_smooth["rel L2(a)"] == pytest.approx(0.3)
    assert physics_smooth["rel L2(u)"] is None
    assert physics_smooth["pde L"] is None
    assert physics_id["rel L2(a)"] is None
    assert str(physics_id["Remark"]).startswith("missing: runs/evaluations/id/")


def test_explicit_smooth_evaluation_takes_priority_over_main_result(tmp_path: Path) -> None:
    main_dir = tmp_path / "runs/main_results/main"
    smooth_dir = (
        tmp_path
        / "runs/evaluations/smooth/task_group=full_forward_main/pde=poisson"
        / "baseline=fno/seed=1/run=eval_smooth_main"
    )
    _write_matrix(
        tmp_path / "matrices/main_results.jsonl",
        [
            {
                "run_id": "main",
                "task_group": "full_forward_main",
                "pde": "poisson",
                "task": "forward",
                "baseline": "fno",
                "seed": 1,
                "output_dir": str(main_dir),
            }
        ],
    )
    _write_result(
        main_dir,
        run_id="main",
        task_group="full_forward_main",
        baseline="fno",
        rel_a=None,
        rel_u=0.1,
        task="forward",
        test_file="poisson/legacy.mat",
    )
    _write_result(
        smooth_dir,
        run_id="eval_smooth_main",
        task_group="full_forward_main",
        baseline="fno",
        rel_a=None,
        rel_u=0.2,
        task="forward",
        test_file="poisson/poisson_test_smooth.mat",
    )

    rows = result_summary.collect_summary_rows(tmp_path)

    smooth = next(row for row in rows if row["DIST"] == "Smooth")
    assert smooth["rel L2(u)"] == pytest.approx(0.2)
    assert smooth["Namespace"] == "evaluations"
    assert smooth["Result Source"] == "explicit_evaluation"
    assert smooth["Fallback Used"] is False
    assert smooth["Test File"] == "poisson/poisson_test_smooth.mat"


def test_burgers_does_not_treat_legacy_main_result_as_smooth(tmp_path: Path) -> None:
    main_dir = tmp_path / "runs/main_results/burger"
    _write_matrix(
        tmp_path / "matrices/main_results.jsonl",
        [
            {
                "run_id": "burger",
                "task_group": "time_varying_da_main",
                "pde": "burger",
                "task": "sparse_solution",
                "baseline": "var4d",
                "seed": 1,
                "output_dir": str(main_dir),
            }
        ],
    )
    _write_result(
        main_dir,
        run_id="burger",
        task_group="time_varying_da_main",
        baseline="var4d",
        rel_a=0.1,
        rel_u=0.2,
        task="sparse_solution",
        pde="burger",
        test_file="burgers/burger_test_10000-128-128.mat",
    )

    rows = result_summary.collect_summary_rows(tmp_path)

    smooth = next(row for row in rows if row["DIST"] == "Smooth")
    assert smooth["rel L2(a)"] == ""
    assert smooth["rel L2(u)"] == ""
    assert smooth["Result Source"] == "missing_evaluation"
    assert smooth["Fallback Used"] is False
    assert smooth["Test File"] == ""


def test_summary_appends_multicondition_ablation_evaluations(tmp_path: Path) -> None:
    main_dir = tmp_path / "runs/main_results/main"
    ablation_dir = (
        tmp_path
        / "runs/evaluations/rough/ablation=sparse_solution_multicondition"
        / "task_group=sparse_solution_multicondition_eval_a_only/pde=poisson"
        / "baseline=recfno/seed=1/run=eval_rough_ablation"
    )
    _write_matrix(
        tmp_path / "matrices/main_results.jsonl",
        [
            {
                "run_id": "main",
                "task_group": "full_forward_main",
                "pde": "poisson",
                "task": "forward",
                "baseline": "fno",
                "seed": 1,
                "output_dir": str(main_dir),
            }
        ],
    )
    _write_result(
        main_dir,
        run_id="main",
        task_group="full_forward_main",
        baseline="fno",
        rel_a=None,
        rel_u=0.1,
        task="forward",
    )
    _write_result(
        ablation_dir,
        run_id="eval_rough_ablation",
        task_group="sparse_solution_multicondition_eval_a_only",
        baseline="recfno",
        rel_a=None,
        rel_u=None,
        task="sparse_solution_multicondition",
        test_file="poisson/poisson_test_rough.mat",
        extra={
            "evaluation_condition_mode": "a_only",
            "rel_l2_a_mean": 0.11,
            "rel_l2_u_mean": 0.22,
            "joint_rel_l2_mean": 0.33,
            "source_train_run_id": "train_ablation",
        },
    )

    rows = result_summary.collect_summary_rows(tmp_path)

    assert len(rows) == 4
    ablation = next(row for row in rows if row["Ablation"])
    assert ablation["DIST"] == "Rough"
    assert ablation["TASK"] == "both"
    assert ablation["CONDITION"] == "a_only"
    assert ablation["rel L2(a)"] == pytest.approx(0.11)
    assert ablation["rel L2(u)"] == pytest.approx(0.22)
    assert ablation["joint rel L2"] == pytest.approx(0.33)
    assert ablation["Namespace"] == "evaluations"
    assert ablation["Result Source"] == "explicit_evaluation"
    assert ablation["Fallback Used"] is False
    assert ablation["Source Run"] == "train_ablation"


def test_summary_uses_out_root_environment_variable(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    run_dir = tmp_path / "runs" / "main_results" / "run"
    _write_matrix(
        tmp_path / "matrices" / "main_results.jsonl",
        [
            {
                "run_id": "run",
                "task_group": "full_forward_main",
                "pde": "poisson",
                "task": "forward",
                "baseline": "fno",
                "seed": 1,
                "output_dir": str(run_dir),
            }
        ],
    )
    _write_result(
        run_dir,
        run_id="run",
        task_group="full_forward_main",
        baseline="fno",
        rel_a=None,
        rel_u=0.125,
        task="forward",
        sensor_mode="none",
    )
    monkeypatch.setenv("OUT_ROOT", str(tmp_path))

    assert result_summary.summary() == tmp_path / "summary" / "results.xlsx"


def test_summary_ignores_quarantined_and_historical_results(tmp_path: Path) -> None:
    current = (
        tmp_path
        / "runs/main_results/task_group=full_forward_main/pde=poisson"
        / "baseline=fno/seed=1/run=current"
    )
    _write_matrix(
        tmp_path / "matrices" / "main_results.jsonl",
        [
            {
                "run_id": "current",
                "task_group": "full_forward_main",
                "pde": "poisson",
                "task": "forward",
                "baseline": "fno",
                "seed": 1,
                "output_dir": str(current),
            }
        ],
    )
    _write_result(
        current,
        run_id="current",
        task_group="full_forward_main",
        baseline="fno",
        rel_a=None,
        rel_u=0.1,
        task="forward",
        sensor_mode="none",
    )
    _write_result(
        current / "quarantine" / "invalid_old",
        run_id="old",
        task_group="full_forward_main",
        baseline="fno",
        rel_a=None,
        rel_u=9.9,
        task="forward",
        sensor_mode="none",
    )
    _write_result(
        current.parent / "run=historical",
        run_id="historical",
        task_group="full_forward_main",
        baseline="fno",
        rel_a=None,
        rel_u=8.8,
        task="forward",
        sensor_mode="none",
    )

    rows = result_summary.collect_summary_rows(tmp_path)

    assert len(rows) == 3
    smooth = next(row for row in rows if row["DIST"] == "Smooth")
    assert smooth["rel L2(u)"] == pytest.approx(0.1)
