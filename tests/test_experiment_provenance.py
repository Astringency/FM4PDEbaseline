from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from baselines.run import main as run_baseline
from baselines.run import parse_args as parse_baseline_args
import scripts.export_results_xlsx as exporter
from scripts.experiments.build_matrix import build_matrix
from scripts.experiments.run_one import build_command, run_one as run_matrix_row
from scripts.experiments.provenance import SUMMARY_DESIGN_FIELD_MAP, repository_revision, run_fingerprint
from scripts.export_results_xlsx import (
    ExportValidationError,
    collect_results,
    export_results,
    require_single_cohort,
)
from scripts.run_experiments import completion_reasons, is_complete, quarantine_invalid_output


ROOT = Path(__file__).resolve().parents[1]


def _config(method_config: Path) -> dict:
    return {
        "name": "provenance_test",
        "experiment_kind": "main",
        "main_table_only": True,
        # These generic provenance unit tests intentionally use a non-paper
        # protocol. Formal v2 manifest enforcement has dedicated coverage in
        # test_data_manifest_provenance.py.
        "task_protocol_version": "provenance-unit-test-v1",
        "config": str(method_config),
        "pdes": ["darcy"],
        "seeds": [1],
        "train_size": 16,
        "val_size": 4,
        "test_size": 8,
        "train_shards": 2,
        "device": "cpu",
        "task_groups": ["full_forward_main"],
        "task_group_overrides": {
            "full_forward_main": {
                "task": "forward",
                "pdes": ["darcy"],
                "baselines": ["fno"],
                "batch_size": 2,
                "epochs": 3,
            }
        },
    }


def _only_row(cfg: dict, tmp_path: Path) -> dict:
    rows, skipped, _summary = build_matrix(cfg, tmp_path / "outputs", "provenance_test")
    assert not skipped
    assert len(rows) == 1
    return rows[0]


def _valid_summary(row: dict) -> dict:
    summary = {
        "status": "success",
        "matrix_schema_version": row["matrix_schema_version"],
        "summary_schema_version": row["summary_schema_version"],
        "run_id": row["run_id"],
        "run_fingerprint": row["run_fingerprint"],
        "execution_mode": row["execution_mode"],
        "eval_only": row["execution_mode"] == "eval_only",
        "comparison_track": row["comparison_track"],
        "config_content_sha256": row["config_content_sha256"],
        "config_hash": "resolved-config-sha1",
        "commit_hash": row["commit_hash"],
        "task_protocol_version": row["task_protocol_version"],
        "sensor_protocol_version": row["sensor_protocol_version"],
        "sensor_seed": row["sensor_seed"],
        "task_group": row["task_group"],
        "task": row["task"],
        "pde": row["pde"],
        "baseline": row["baseline"],
        "seed": row["seed"],
    }
    summary.update(
        {
            summary_field: row[matrix_field]
            for matrix_field, summary_field in SUMMARY_DESIGN_FIELD_MAP.items()
        }
    )
    summary["train_size_requested"] = row["train_size"]
    summary["val_size"] = row["val_size"]
    summary["test_size"] = row["test_size"]
    return summary


def test_matrix_fingerprint_tracks_config_content(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    cfg = _config(method_config)

    first = _only_row(cfg, tmp_path / "first")
    expected_config_hash = hashlib.sha256(method_config.read_bytes()).hexdigest()

    assert first["matrix_schema_version"] == 2
    assert first["summary_schema_version"] == 2
    assert first["execution_mode"] == "train"
    assert first["config_content_sha256"] == expected_config_hash
    assert first["task_protocol_version"]
    assert first["sensor_protocol_version"]
    assert first["commit_hash"] == repository_revision(Path(__file__).resolve().parents[1])
    assert len(first["run_fingerprint"]) == 64
    assert first["run_id"].endswith(first["run_fingerprint"][:12])

    method_config.write_text("method:\n  width: 32\n", encoding="utf-8")
    changed = _only_row(cfg, tmp_path / "changed")

    assert changed["config_content_sha256"] != first["config_content_sha256"]
    assert changed["run_fingerprint"] != first["run_fingerprint"]
    assert changed["run_id"] != first["run_id"]


def test_matrix_rejects_nonignored_in_repository_output_root(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")

    with pytest.raises(ValueError, match="not git-ignored"):
        build_matrix(
            _config(method_config),
            ROOT / "nonignored-matrix-output-fixture",
            "unsafe",
        )


def test_repository_revision_covers_tracked_and_untracked_content(tmp_path: Path):
    repo = tmp_path / "repo"
    repo.mkdir()
    subprocess.run(["git", "init", "-q"], cwd=repo, check=True)
    subprocess.run(["git", "config", "user.name", "Test User"], cwd=repo, check=True)
    subprocess.run(["git", "config", "user.email", "test@example.invalid"], cwd=repo, check=True)
    tracked = repo / "tracked.txt"
    tracked.write_text("one\n", encoding="utf-8")
    subprocess.run(["git", "add", "tracked.txt"], cwd=repo, check=True)
    subprocess.run(["git", "commit", "-qm", "initial"], cwd=repo, check=True)

    head = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=repo, text=True).strip()
    assert repository_revision(repo) == head

    tracked.write_text("two\n", encoding="utf-8")
    tracked_revision = repository_revision(repo)
    assert tracked_revision.startswith(f"{head}-dirty-")
    assert repository_revision(repo) == tracked_revision

    untracked = repo / "untracked.bin"
    untracked.write_bytes(b"first")
    first_untracked_revision = repository_revision(repo)
    assert first_untracked_revision != tracked_revision
    untracked.write_bytes(b"second")
    assert repository_revision(repo) != first_untracked_revision


def test_run_fingerprint_tracks_repository_revision(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    changed = dict(row)
    changed["commit_hash"] = f"{row['commit_hash']}-different"

    assert run_fingerprint(changed) != row["run_fingerprint"]


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("val_size", 5),
        ("train_shards", 3),
        ("device", "cuda"),
        ("sensor_seed", 17),
        ("task_protocol_version", "task-contract-test-v99"),
        ("sensor_protocol_version", "sensor-contract-test-v99"),
    ],
)
def test_matrix_fingerprint_tracks_protocol_and_execution_fields(
    tmp_path: Path,
    field: str,
    value: object,
):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    base_cfg = _config(method_config)
    changed_cfg = deepcopy(base_cfg)
    changed_cfg[field] = value

    first = _only_row(base_cfg, tmp_path / "first")
    changed = _only_row(changed_cfg, tmp_path / "changed")

    assert changed["run_fingerprint"] != first["run_fingerprint"]
    assert changed["run_id"] != first["run_id"]


@pytest.mark.parametrize(("field", "value"), [("batch_size", 4), ("epochs", 4)])
def test_matrix_fingerprint_tracks_group_training_fields(
    tmp_path: Path,
    field: str,
    value: int,
):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    base_cfg = _config(method_config)
    changed_cfg = deepcopy(base_cfg)
    changed_cfg["task_group_overrides"]["full_forward_main"][field] = value

    first = _only_row(base_cfg, tmp_path / "first")
    changed = _only_row(changed_cfg, tmp_path / "changed")

    assert changed["run_fingerprint"] != first["run_fingerprint"]
    assert changed["run_id"] != first["run_id"]


def test_remaining_runner_accepts_only_a_successful_matching_v2_summary(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)

    assert not is_complete(row)
    assert completion_reasons(row) == ["summary_missing"]

    summary_path = output_dir / "summary.json"
    summary_path.write_text(json.dumps({"run_id": row["run_id"]}), encoding="utf-8")
    assert not is_complete(row)
    assert "legacy_matrix_missing:run_fingerprint" not in completion_reasons(row)
    assert "summary_missing:run_fingerprint" in completion_reasons(row)

    valid = _valid_summary(row)
    summary_path.write_text(json.dumps(valid), encoding="utf-8")
    assert is_complete(row)
    assert completion_reasons(row) == []

    for field, bad_value in [
        ("status", "failed"),
        ("run_id", "another-run"),
        ("run_fingerprint", "0" * 64),
        ("config_content_sha256", "1" * 64),
        ("commit_hash", "different-repository-revision"),
        ("execution_mode", "eval_only"),
    ]:
        invalid = dict(valid)
        invalid[field] = bad_value
        summary_path.write_text(json.dumps(invalid), encoding="utf-8")
        assert not is_complete(row), field


def test_completion_rejects_an_internally_inconsistent_matrix_fingerprint(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    row["epochs"] += 1
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps(_valid_summary(row)), encoding="utf-8")

    assert not is_complete(row)
    assert "matrix_mismatch:run_fingerprint" in completion_reasons(row)


def test_run_command_forwards_matrix_provenance_to_core_runner(tmp_path: Path, monkeypatch):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "PDEdata"))

    command = build_command(row)

    expected = {
        "--summary-schema-version": str(row["summary_schema_version"]),
        "--run-fingerprint": row["run_fingerprint"],
        "--config-content-sha256": row["config_content_sha256"],
        "--task-protocol-version": row["task_protocol_version"],
        "--sensor-protocol-version": row["sensor_protocol_version"],
        "--execution-mode": row["execution_mode"],
        "--commit-hash": row["commit_hash"],
    }
    for flag, value in expected.items():
        index = command.index(flag)
        assert command[index + 1] == value

    parsed = parse_baseline_args(command[3:])
    assert parsed.run_fingerprint == row["run_fingerprint"]
    assert parsed.config_content_sha256 == row["config_content_sha256"]
    assert parsed.summary_schema_version == row["summary_schema_version"]
    assert parsed.execution_mode == "train"
    assert parsed.task_protocol_version == row["task_protocol_version"]
    assert parsed.sensor_protocol_version == row["sensor_protocol_version"]
    assert parsed.comparison_track == "unified_adapted"
    assert parsed.commit_hash == row["commit_hash"]


def test_core_runner_rejects_supplied_repository_revision_mismatch():
    with pytest.raises(ValueError, match="does not match the current repository revision"):
        run_baseline(
            [
                "--baseline",
                "fno",
                "--pde",
                "poisson",
                "--commit-hash",
                "definitely-not-the-current-revision",
            ]
        )


def test_run_command_rejects_config_drift_after_matrix_generation(tmp_path: Path, monkeypatch):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "PDEdata"))

    method_config.write_text("method:\n  width: 32\n", encoding="utf-8")

    with pytest.raises(RuntimeError, match="matrix provenance mismatch"):
        build_command(row)


def test_core_runner_recomputes_supplied_fingerprint_before_data_loading():
    with pytest.raises(ValueError, match="effective execution design"):
        run_baseline(
            [
                "--baseline",
                "fno",
                "--pde",
                "poisson",
                "--experiment-mode",
                "debug",
                "--synthetic-data",
                "--run-fingerprint",
                "0" * 64,
            ]
        )


def test_run_command_rejects_fingerprint_bound_environment_override(tmp_path: Path, monkeypatch):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "PDEdata"))
    monkeypatch.setenv("ALLOW_ROW_OVERRIDE", "1")
    monkeypatch.setenv("EPOCHS", "99")

    with pytest.raises(RuntimeError, match="regenerate the matrix"):
        build_command(row)


def test_exporter_defaults_to_the_corrected_namespace():
    args = exporter.parse_args([])

    assert Path(args.matrix) == (
        exporter.OUTPUT_ROOT
        / "experiment_plan_v2_corrected"
        / "matrices"
        / "experiment_plan_v2_corrected.jsonl"
    )
    assert Path(args.output) == exporter.OUTPUT_ROOT / "experiment_plan_v2_corrected_summary.xlsx"
    assert exporter.IMMUTABLE_HISTORICAL_WORKBOOK == (
        exporter.ROOT / "outputs" / "experiment_plan_v2_summary.xlsx"
    )
    assert Path(args.output) != exporter.IMMUTABLE_HISTORICAL_WORKBOOK


def test_exporter_quarantines_legacy_summary_and_keeps_provenance_fields(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    cfg = _config(method_config)
    cfg["seeds"] = [1, 2]
    rows, skipped, _summary = build_matrix(cfg, tmp_path / "matrix", "provenance_test")
    assert not skipped
    assert len(rows) == 2

    good = _valid_summary(rows[0]) | {
        "checkpoint_path": "/checkpoints/model.pt",
        "checkpoint_sha256": "a" * 64,
        "file_paths_summary": '{"train": {"count": 2}}',
        "train_time": 12.5,
        "inference_time_total": 3.0,
        "inference_time_per_sample": 0.3,
        "inference_optimization_time_total": 2.0,
        "inference_optimization_time_per_sample": 0.2,
        "relative_l2_solution_mean": 0.1,
    }
    for row, summary in [(rows[0], good), (rows[1], {"run_id": rows[1]["run_id"]})]:
        output_dir = Path(row["output_dir"])
        output_dir.mkdir(parents=True)
        (output_dir / "summary.json").write_text(json.dumps(summary), encoding="utf-8")

    result = collect_results(rows)

    assert len(result.records) == 1
    assert len(result.quarantine) == 1
    record = result.records[0]
    assert record["publishable"] is True
    assert record["execution_mode"] == "train"
    assert record["run_fingerprint"] == rows[0]["run_fingerprint"]
    assert record["config_content_sha256"] == rows[0]["config_content_sha256"]
    assert record["checkpoint_path"] == "/checkpoints/model.pt"
    assert record["inference_optimization_time_total"] == 2.0
    assert "summary_missing:run_fingerprint" in result.quarantine[0]["validation_reasons"]
    assert result.quarantine[0]["summary_path"].endswith("summary.json")


def test_exporter_quarantines_every_mismatched_design_field(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    summary_path = output_dir / "summary.json"
    valid = _valid_summary(row)

    for matrix_field, summary_field in SUMMARY_DESIGN_FIELD_MAP.items():
        expected = row[matrix_field]
        if isinstance(expected, bool):
            mismatched_value = not expected
        elif isinstance(expected, (int, float)):
            mismatched_value = expected + 1
        else:
            mismatched_value = f"{expected}-mismatch"
        summary_path.write_text(
            json.dumps(valid | {summary_field: mismatched_value}),
            encoding="utf-8",
        )

        result = collect_results([row])

        assert not result.records, summary_field
        assert len(result.quarantine) == 1, summary_field
        assert f"summary_mismatch:{summary_field}" in result.quarantine[0]["validation_reasons"]


def test_exporter_rejects_multiple_provenance_cohorts(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    cfg = _config(method_config)
    cfg["seeds"] = [1, 2]
    rows, _skipped, _summary = build_matrix(cfg, tmp_path / "matrix", "provenance_test")

    rows[1]["commit_hash"] = f"{rows[1]['commit_hash']}-alternate-cohort"
    rows[1]["run_fingerprint"] = run_fingerprint(rows[1])
    for row in rows:
        summary = _valid_summary(row)
        output_dir = Path(row["output_dir"])
        output_dir.mkdir(parents=True)
        (output_dir / "summary.json").write_text(json.dumps(summary), encoding="utf-8")

    result = collect_results(rows)
    assert not result.quarantine
    with pytest.raises(ExportValidationError, match="multiple provenance cohorts"):
        require_single_cohort(result.records)


def test_exporter_writes_provenance_and_timing_columns_for_one_valid_cohort(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    summary = _valid_summary(row) | {
        "checkpoint_path": "/checkpoints/model.pt",
        "checkpoint_sha256": "a" * 64,
        "train_time": 12.5,
        "inference_time_total": 3.0,
        "inference_optimization_time_total": 2.0,
        "relative_l2_solution_mean": 0.1,
    }
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps(summary), encoding="utf-8")
    output = tmp_path / "validated.xlsx"

    report = export_results(
        [row],
        output=output,
        quarantine_output=tmp_path / "quarantine.xlsx",
        quarantine_jsonl=tmp_path / "quarantine.jsonl",
        matrix_path=tmp_path / "matrix.jsonl",
    )

    assert report["publication_blocked"] is False
    workbook = load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == ["runs", "manifest"]
    rows = list(workbook["runs"].iter_rows(values_only=True))
    headers = list(rows[0])
    values = dict(zip(headers, rows[1], strict=True))
    assert values["publishable"] is True
    assert values["execution_mode"] == "train"
    assert values["run_fingerprint"] == row["run_fingerprint"]
    assert values["checkpoint_path"] == "/checkpoints/model.pt"
    assert values["inference_optimization_time_total"] == 2.0


@pytest.mark.parametrize(
    ("checkpoint_state", "expected_reason"),
    [
        ("missing", "checkpoint_missing"),
        ("tampered", "checkpoint_hash_mismatch"),
    ],
)
def test_exporter_quarantines_eval_only_summary_when_live_checkpoint_is_invalid(
    tmp_path: Path,
    checkpoint_state: str,
    expected_reason: str,
):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    train_row = _only_row(_config(method_config), tmp_path / "matrix")
    checkpoint = tmp_path / "model.pt"
    checkpoint.write_bytes(b"checkpoint-at-evaluation-time")
    eval_row = deepcopy(train_row)
    eval_row.update(
        {
            "execution_mode": "eval_only",
            "source_train_run_id": train_row["run_id"],
            "source_train_run_fingerprint": train_row["run_fingerprint"],
            "source_train_seed": train_row["seed"],
            "checkpoint_path": str(checkpoint),
            "checkpoint_sha256": hashlib.sha256(checkpoint.read_bytes()).hexdigest(),
        }
    )
    eval_row["run_fingerprint"] = run_fingerprint(eval_row)
    eval_row["run_id"] = f"eval_{eval_row['run_fingerprint'][:12]}"
    summary = _valid_summary(eval_row) | {
        "source_train_run_id": eval_row["source_train_run_id"],
        "source_train_run_fingerprint": eval_row["source_train_run_fingerprint"],
        "source_train_seed": eval_row["source_train_seed"],
        "checkpoint_path": eval_row["checkpoint_path"],
        "checkpoint_sha256": eval_row["checkpoint_sha256"],
    }
    output_dir = Path(eval_row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps(summary), encoding="utf-8")
    valid_collection = collect_results([eval_row])
    assert len(valid_collection.records) == 1
    assert not valid_collection.quarantine
    if checkpoint_state == "missing":
        checkpoint.unlink()
    else:
        checkpoint.write_bytes(b"checkpoint-tampered-after-summary")
    output = tmp_path / "must_not_publish_eval.xlsx"
    quarantine = tmp_path / "eval_quarantine.xlsx"
    quarantine_jsonl = tmp_path / "eval_quarantine.jsonl"

    with pytest.raises(ExportValidationError, match="publication blocked"):
        export_results(
            [eval_row],
            output=output,
            quarantine_output=quarantine,
            quarantine_jsonl=quarantine_jsonl,
        )

    assert not output.exists()
    assert quarantine.exists()
    quarantined = json.loads(quarantine_jsonl.read_text(encoding="utf-8"))
    assert expected_reason in quarantined["validation_reasons"]


def test_exporter_blocks_main_workbook_and_writes_readable_quarantine(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps({"run_id": row["run_id"]}), encoding="utf-8")
    output = tmp_path / "must_not_publish.xlsx"
    quarantine = tmp_path / "quarantine.xlsx"
    quarantine_jsonl = tmp_path / "quarantine.jsonl"
    output.write_bytes(b"historical workbook that must not remain published")

    with pytest.raises(ExportValidationError, match="publication blocked"):
        export_results(
            [row],
            output=output,
            quarantine_output=quarantine,
            quarantine_jsonl=quarantine_jsonl,
        )

    assert not output.exists()
    assert quarantine.exists()
    assert quarantine_jsonl.exists()
    archived = list(tmp_path.glob("quarantine_previous_publication*.xlsx"))
    assert len(archived) == 1
    assert archived[0].read_bytes() == b"historical workbook that must not remain published"
    workbook = load_workbook(quarantine, read_only=True, data_only=True)
    assert workbook.sheetnames == ["quarantine", "validated_preview", "manifest"]
    quarantine_rows = list(workbook["quarantine"].iter_rows(values_only=True))
    headers = list(quarantine_rows[0])
    values = dict(zip(headers, quarantine_rows[1], strict=True))
    assert values["publishable"] is False
    assert "summary_missing:run_fingerprint" in values["validation_reasons"]


def test_exporter_preserves_immutable_historical_workbook_when_publication_is_blocked(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps({"run_id": row["run_id"]}), encoding="utf-8")
    historical = tmp_path / "experiment_plan_v2_summary.xlsx"
    historical_bytes = b"immutable historical workbook"
    historical.write_bytes(historical_bytes)
    monkeypatch.setattr(exporter, "IMMUTABLE_HISTORICAL_WORKBOOK", historical)
    quarantine = tmp_path / "quarantine.xlsx"
    quarantine_jsonl = tmp_path / "quarantine.jsonl"

    with pytest.raises(ExportValidationError, match="publication blocked"):
        export_results(
            [row],
            output=historical,
            quarantine_output=quarantine,
            quarantine_jsonl=quarantine_jsonl,
        )

    assert historical.read_bytes() == historical_bytes
    assert quarantine.exists()
    assert quarantine_jsonl.exists()
    assert not list(tmp_path.glob("quarantine_previous_publication*.xlsx"))
    manifest_rows = list(
        load_workbook(quarantine, read_only=True, data_only=True)["manifest"].iter_rows(values_only=True)
    )
    manifest = dict(zip(manifest_rows[0], manifest_rows[1], strict=True))
    assert manifest["immutable_historical_workbook_preserved"] == str(historical)


def test_exporter_refuses_to_overwrite_immutable_historical_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps(_valid_summary(row)), encoding="utf-8")
    historical = tmp_path / "experiment_plan_v2_summary.xlsx"
    historical_bytes = b"immutable historical workbook"
    historical.write_bytes(historical_bytes)
    monkeypatch.setattr(exporter, "IMMUTABLE_HISTORICAL_WORKBOOK", historical)

    with pytest.raises(ExportValidationError, match="immutable historical workbook"):
        export_results(
            [row],
            output=historical,
            quarantine_output=tmp_path / "quarantine.xlsx",
            quarantine_jsonl=tmp_path / "quarantine.jsonl",
        )

    assert historical.read_bytes() == historical_bytes


@pytest.mark.parametrize("historical_target", ["quarantine_output", "quarantine_jsonl"])
def test_exporter_refuses_to_use_historical_workbook_as_a_quarantine_target(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    historical_target: str,
):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    historical = tmp_path / "experiment_plan_v2_summary.xlsx"
    historical_bytes = b"immutable historical workbook"
    historical.write_bytes(historical_bytes)
    monkeypatch.setattr(exporter, "IMMUTABLE_HISTORICAL_WORKBOOK", historical)
    targets = {
        "output": tmp_path / "corrected.xlsx",
        "quarantine_output": tmp_path / "quarantine.xlsx",
        "quarantine_jsonl": tmp_path / "quarantine.jsonl",
    }
    targets[historical_target] = historical

    with pytest.raises(ExportValidationError, match="immutable historical workbook"):
        export_results([row], **targets)

    assert historical.read_bytes() == historical_bytes


def test_exporter_blocks_partial_matrix_when_a_summary_is_missing(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    cfg = _config(method_config)
    cfg["seeds"] = [1, 2]
    rows, _skipped, _summary = build_matrix(cfg, tmp_path / "matrix", "provenance_test")
    first_output = Path(rows[0]["output_dir"])
    first_output.mkdir(parents=True)
    (first_output / "summary.json").write_text(json.dumps(_valid_summary(rows[0])), encoding="utf-8")
    output = tmp_path / "must_not_publish_partial.xlsx"

    with pytest.raises(ExportValidationError, match="publication blocked"):
        export_results(
            rows,
            output=output,
            quarantine_output=tmp_path / "partial_quarantine.xlsx",
            quarantine_jsonl=tmp_path / "partial_quarantine.jsonl",
        )

    assert not output.exists()
    collection = collect_results(rows)
    assert collection.missing_run_ids == [rows[1]["run_id"]]
    assert any(item["validation_reasons"] == ["summary_missing"] for item in collection.quarantine)


def test_exporter_exposes_execution_and_comparison_track_audit_fields(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    summary = _valid_summary(row) | {"eval_only": False}
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps(summary), encoding="utf-8")

    record = collect_results([row]).records[0]

    assert record["status"] == "success"
    assert record["eval_only"] is False
    assert record["official_native_eligible"] is row["official_native_eligible"]


def test_comparison_tracks_separate_unified_adapted_from_official_native(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    base = {
        "name": "track_test",
        "experiment_kind": "main",
        "main_table_only": False,
        "config": str(method_config),
        "pdes": ["darcy"],
        "seeds": [1],
        "train_size": 16,
        "val_size": 4,
        "test_size": 8,
        "train_shards": 2,
        "device": "cpu",
        "task_groups": ["sparse_solution_main_amortized"],
        "task_group_overrides": {
            "sparse_solution_main_amortized": {
                "task": "sparse_solution",
                "pdes": ["darcy"],
                "baselines": ["recfno"],
                "num_sensors": 8,
                "batch_size": 2,
                "epochs": 3,
            }
        },
    }
    unified = deepcopy(base) | {"comparison_track": "unified_adapted"}
    official = deepcopy(base) | {"comparison_track": "official_native"}

    unified_rows, unified_skips, _ = build_matrix(unified, tmp_path / "unified", "unified")
    official_rows, official_skips, _ = build_matrix(official, tmp_path / "official", "official")

    assert not unified_skips
    assert len(unified_rows) == 1
    assert unified_rows[0]["comparison_track"] == "unified_adapted"
    assert unified_rows[0]["sensor_mode"] == "random_per_sample"
    assert unified_rows[0]["capability_status"] == "adapted"
    assert unified_rows[0]["implementation_required"] == "adapted_allowed"
    assert unified_rows[0]["official_native_eligible"] is False
    assert not official_rows
    assert len(official_skips) == 1
    assert "official-native" in official_skips[0]["reason"]


def test_remaining_runner_preserves_invalid_artifacts_in_quarantine_before_rerun(tmp_path: Path):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps({"run_id": row["run_id"]}), encoding="utf-8")
    (output_dir / "results_raw.jsonl").write_text('{"old": true}\n', encoding="utf-8")
    (output_dir / "results_summary.jsonl").write_text('{"old": true}\n', encoding="utf-8")

    quarantine_dir = quarantine_invalid_output(row)

    assert quarantine_dir is not None
    assert not (output_dir / "summary.json").exists()
    assert not (output_dir / "results_raw.jsonl").exists()
    assert (quarantine_dir / "summary.json").exists()
    assert (quarantine_dir / "results_raw.jsonl").exists()
    manifest = json.loads((quarantine_dir / "quarantine_manifest.json").read_text(encoding="utf-8"))
    assert manifest["run_id"] == row["run_id"]
    assert "summary_missing:run_fingerprint" in manifest["validation_reasons"]


def test_matrix_runner_skips_only_when_matching_summary_is_valid(tmp_path: Path, monkeypatch):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps(_valid_summary(row)), encoding="utf-8")

    def unexpected_launch(*_args, **_kwargs):
        raise AssertionError("valid completed run must not launch a subprocess")

    monkeypatch.setattr("scripts.experiments.run_one.subprocess.Popen", unexpected_launch)

    assert run_matrix_row(row, ["must-not-run"]) == 0
    assert (output_dir / "run.done").exists()


def test_matrix_runner_quarantines_stale_done_artifacts_before_relaunch(tmp_path: Path, monkeypatch):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    output_dir.mkdir(parents=True)
    (output_dir / "summary.json").write_text(json.dumps({"run_id": row["run_id"]}), encoding="utf-8")
    (output_dir / "run.done").write_text('{"legacy": true}', encoding="utf-8")
    valid_json = json.dumps(_valid_summary(row))
    code = (
        "from pathlib import Path; "
        f"Path({str(output_dir / 'summary.json')!r}).write_text({valid_json!r}, encoding='utf-8')"
    )
    monkeypatch.setenv("PROGRESS_INTERVAL_SECONDS", "0")

    assert run_matrix_row(row, [sys.executable, "-c", code]) == 0

    quarantine_dirs = list((output_dir / "quarantine").glob("invalid_*"))
    assert len(quarantine_dirs) == 1
    assert json.loads((quarantine_dirs[0] / "summary.json").read_text(encoding="utf-8")) == {
        "run_id": row["run_id"]
    }
    assert (quarantine_dirs[0] / "run.done").exists()
    assert is_complete(row)


def test_matrix_runner_does_not_mark_exit_zero_without_valid_summary_done(tmp_path: Path, monkeypatch):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    row = _only_row(_config(method_config), tmp_path / "matrix")
    output_dir = Path(row["output_dir"])
    monkeypatch.setenv("PROGRESS_INTERVAL_SECONDS", "0")

    assert run_matrix_row(row, [sys.executable, "-c", "pass"]) == 1
    assert not (output_dir / "run.done").exists()
    failure = json.loads((output_dir / "run.failed").read_text(encoding="utf-8"))
    assert failure["exit_code"] == 1
    assert failure["summary_validation_reasons"] == ["summary_missing"]


def test_eval_only_has_its_own_fingerprint_and_links_verified_training_checkpoint(tmp_path: Path, monkeypatch):
    method_config = tmp_path / "method.yaml"
    method_config.write_text("method:\n  width: 16\n", encoding="utf-8")
    train_row = _only_row(_config(method_config), tmp_path / "matrix")
    checkpoint = tmp_path / "model.pt"
    checkpoint.write_bytes(b"checkpoint-for-provenance-test")
    eval_row = deepcopy(train_row)
    eval_row.update(
        {
            "execution_mode": "eval_only",
            "source_train_run_id": train_row["run_id"],
            "source_train_run_fingerprint": train_row["run_fingerprint"],
            "source_train_seed": train_row["seed"],
            "sensor_seed": train_row["sensor_seed"] + 1,
            "checkpoint_path": str(checkpoint),
            "checkpoint_sha256": hashlib.sha256(checkpoint.read_bytes()).hexdigest(),
        }
    )
    eval_row["run_fingerprint"] = run_fingerprint(eval_row)
    eval_row["run_id"] = f"eval_{eval_row['run_fingerprint'][:12]}"
    monkeypatch.setenv("DATA_ROOT", str(tmp_path / "PDEdata"))

    assert eval_row["run_fingerprint"] != train_row["run_fingerprint"]
    summary = _valid_summary(eval_row) | {
        "source_train_run_id": train_row["run_id"],
        "source_train_run_fingerprint": train_row["run_fingerprint"],
        "source_train_seed": train_row["seed"],
        "checkpoint_path": eval_row["checkpoint_path"],
        "checkpoint_sha256": eval_row["checkpoint_sha256"],
    }
    assert completion_reasons(
        eval_row | {"output_dir": str(tmp_path / "not-written-yet")}
    ) == ["summary_missing"]
    output_dir = Path(eval_row["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "summary.json").write_text(json.dumps(summary), encoding="utf-8")
    assert is_complete(eval_row)

    wrong_source_seed = dict(summary, source_train_seed=train_row["seed"] + 1)
    (output_dir / "summary.json").write_text(json.dumps(wrong_source_seed), encoding="utf-8")
    assert not is_complete(eval_row)
    assert "summary_mismatch:source_train_seed" in completion_reasons(eval_row)
    (output_dir / "summary.json").write_text(json.dumps(summary), encoding="utf-8")

    command = build_command(eval_row)
    assert "--eval-only" in command
    assert command[command.index("--checkpoint") + 1] == eval_row["checkpoint_path"]
    assert command[command.index("--source-train-run-id") + 1] == train_row["run_id"]
    assert command[command.index("--source-train-run-fingerprint") + 1] == train_row["run_fingerprint"]
    assert command[command.index("--source-train-seed") + 1] == str(train_row["seed"])
    assert command[command.index("--sensor-seed") + 1] == str(eval_row["sensor_seed"])
